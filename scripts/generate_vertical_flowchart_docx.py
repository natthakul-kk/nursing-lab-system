import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_vertical_flowchart_docx():
    target_docx = r"d:\LAB-system\system_flow\ผังขั้นตอนการทำงานทั้งระบบ_System_Flowchart.docx"
    alias_docx = r"d:\LAB-system\system_flow\system_operational_flowchart.docx"
    png_path = r"d:\LAB-system\system_flow\ผังขั้นตอนการทำงานทั้งระบบ_System_Flowchart.png"

    doc = docx.Document()

    # Page Margins (Normal 1 inch or 0.8 inch)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Set default font to TH Sarabun New 16pt
    style = doc.styles['Normal']
    font = style.font
    font.name = 'TH Sarabun New'
    font.size = Pt(16)
    font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)

    def set_font_run(run, name='TH Sarabun New', size=16, bold=False, italic=False, color=RGBColor(0x0f, 0x17, 0x2a)):
        run.font.name = name
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        # Ensure complex script font is also set
        rPr = run._r.get_or_add_rPr()
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{name}" w:hAnsi="{name}" w:cs="{name}"/>')
        rPr.append(rFonts)

    def add_heading_styled(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        if level == 1:
            set_font_run(run, 'TH Sarabun New', size=22, bold=True, color=RGBColor(0x02, 0x84, 0xc7))
        elif level == 2:
            set_font_run(run, 'TH Sarabun New', size=18, bold=True, color=RGBColor(0x0f, 0x17, 0x2a))
        elif level == 3:
            set_font_run(run, 'TH Sarabun New', size=16, bold=True, color=RGBColor(0x33, 0x41, 0x55))
        return p

    def add_body_p(text, bold_prefix="", bullet=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bullet:
            p.paragraph_format.left_indent = Inches(0.25)
            r_bullet = p.add_run("• ")
            set_font_run(r_bullet, 'TH Sarabun New', size=16, bold=True, color=RGBColor(0x02, 0x84, 0xc7))
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            set_font_run(r_bold, 'TH Sarabun New', size=16, bold=True)
        r_text = p.add_run(text)
        set_font_run(r_text, 'TH Sarabun New', size=16)
        return p

    # --- Title Section ---
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(2)
    r_badge = p_title.add_run("เอกสารผังขั้นตอนการปฏิบัติงาน (STANDARD OPERATIONAL FLOWCHART)\n")
    set_font_run(r_badge, 'TH Sarabun New', size=14, bold=True, color=RGBColor(0x02, 0x84, 0xc7))
    r_main_title = p_title.add_run("ผังขั้นตอนการทำงานและความเชื่อมโยงทั้งระบบ (Vertical Flowchart)\n")
    set_font_run(r_main_title, 'TH Sarabun New', size=24, bold=True, color=RGBColor(0x0f, 0x17, 0x2a))
    r_sub = p_title.add_run("ระบบห้องปฏิบัติการพยาบาล (Nursing Laboratory Information Management System)")
    set_font_run(r_sub, 'TH Sarabun New', size=16, italic=True, color=RGBColor(0x64, 0x74, 0x8b))

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- Embed Vertical Flowchart Image ---
    if os.path.exists(png_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        # Scale image to fit width (approx 6.5 inches)
        run_img.add_picture(png_path, width=Inches(6.5))

        p_caption = doc.add_paragraph()
        p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_caption.paragraph_format.space_after = Pt(14)
        r_cap = p_caption.add_run("รูปที่ 1: แผนผังขั้นตอนการทำงานทั้งระบบ (System Operational Flowchart)")
        set_font_run(r_cap, 'TH Sarabun New', size=14, italic=True, color=RGBColor(0x64, 0x74, 0x8b))

    # --- Overview Narrative ---
    add_heading_styled("1. บทนำและโครงสร้างภาพรวมของผังงาน (Overview)", level=1)
    add_body_p(
        "ผังงานขั้นตอนการทำงานและความเชื่อมโยงทั้งระบบ (Vertical Flowchart) ได้รับการออกแบบตามมาตรฐานสากล โดยเรียงลำดับขั้นตอนจากบนลงล่าง (Top-to-Bottom) เพื่อแสดงวงจรชีวิตของกระบวนการทำงานจริง (End-to-End Operational Lifecycle) ภายในระบบห้องปฏิบัติการพยาบาล ตั้งแต่การจัดเตรียมรายวิชาและคลังพัสดุ การยื่นคำร้อง การอนุมัติ การจ่าย-คืน จนถึงการตรวจรับความเสียหายและการประมวลผลดัชนีชี้วัดสู่ระดับผู้บริหาร"
    )

    # --- Symbology Table ---
    add_heading_styled("2. สัญลักษณ์และความหมายที่ใช้ในผังงาน (Flowchart Symbology)", level=2)
    sym_table = doc.add_table(rows=1, cols=3)
    sym_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = sym_table.rows[0].cells
    hdr_titles = ["สัญลักษณ์ / รูปทรง", "ความหมายตามมาตรฐาน", "การนำมาใช้ในระบบ"]
    for i, title in enumerate(hdr_titles):
        hdr_cells[i].text = title
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.name = 'TH Sarabun New'
        p.runs[0].font.size = Pt(16)
        p.runs[0].font.bold = True
        set_font_run(p.runs[0], 'TH Sarabun New', size=16, bold=True, color=RGBColor(0xff, 0xff, 0xff))
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0284c7"/>')
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)

    sym_data = [
        ("Terminator (วงรี/แคปซูล)", "จุดเริ่มต้นหรือสิ้นสุดกระบวนการ", "จุดเข้าสู่ระบบ และจุดสิ้นสุดการปิดคำร้องสมบูรณ์"),
        ("Process (สี่เหลี่ยมผืนผ้า)", "ขั้นตอนการปฏิบัติงาน/การประมวลผล", "การยื่นคำร้อง, จัดของ, สแกน QR Code, การคืนของ"),
        ("Decision Diamond (ข้าวหลามตัด)", "จุดตัดสินใจและเงื่อนไข (มีทางแยก ใช่/ไม่ใช่)", "การตรวจสอบสต็อก, การพิจารณาอนุมัติ, การตรวจความเสียหาย"),
        ("Data / Storage (ฐานข้อมูล)", "การจัดเก็บหรือดึงข้อมูลจากระบบ", "การบันทึกสถานะคำร้อง, Audit Trail, สถิติ KPIs")
    ]

    for row_data in sym_data:
        row = sym_table.add_row()
        for idx, text in enumerate(row_data):
            cell = row.cells[idx]
            cell.text = text
            p = cell.paragraphs[0]
            set_font_run(p.runs[0], 'TH Sarabun New', size=16)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- 6 Operational Stages Detailed ---
    add_heading_styled("3. รายละเอียดขั้นตอนการดำเนินงาน 6 ขั้นตอนหลัก (Detailed Operational Steps)", level=1)

    stages = [
        (
            "ขั้นตอนที่ 1: การเตรียมข้อมูลแม่แบบและบริหารคลังพัสดุ (Master Data & Inventory Setup)",
            "การเตรียมความพร้อมก่อนเปิดภาคการศึกษา เพื่อให้การเบิกจ่ายและการจองฝึกปฏิบัติเป็นไปอย่างราบรื่น",
            [
                ("การสร้างรายวิชาและชุดแม่แบบ: ", "อาจารย์หรือเจ้าหน้าที่บันทึกรายวิชาและผูกเกณฑ์การจัดอุปกรณ์มาตรฐาน (Kit Templates)"),
                ("การตรวจสอบยอดคงคลัง: ", "ระบบตรวจสอบปริมาณเวชภัณฑ์ (Stock Lots) และสถานะความพร้อมของครุภัณฑ์ (Equipment Assets)"),
                ("จุดตัดสินใจที่ 1 (ยอดพัสดุเพียงพอหรือไม่?): ", "หากไม่พอ ระบบจะแจ้งเตือน Low Stock เพื่อเปิดกระบวนการจัดซื้อและตรวจรับพัสดุใหม่ (Stock-In) หากเพียงพอ จะนำไปสู่การแบ่งบรรจุซองย่อยสเตอร์ไรด์ (Repack) และประกอบชุดฝึกพร้อมใช้")
            ]
        ),
        (
            "ขั้นตอนที่ 2: การยื่นคำร้องและการจองฝึกปฏิบัติ (Request & Booking Submission)",
            "ขั้นตอนที่นักศึกษาหรืออาจารย์ผู้สอนทำการขอใช้บริการห้องปฏิบัติการและเบิกพัสดุ",
            [
                ("การระบุรายละเอียดคำร้อง: ", "ผู้ขอเลือกรายวิชา กำหนดวัน-เวลาที่ต้องการใช้ห้อง และระบุชุดอุปกรณ์หรือพัสดุที่ต้องการ"),
                ("การส่งคำร้องเข้าสู่ระบบ: ", "ระบบบันทึกคำร้องลงฐานข้อมูลกลาง พร้อมกำหนดสถานะเริ่มต้นเป็น 'รอดำเนินการ' (PENDING)")
            ]
        ),
        (
            "ขั้นตอนที่ 3: การตรวจสอบและพิจารณาอนุมัติคำขอ (Review & Approval Process)",
            "กลไกการกลั่นกรองและควบคุมการใช้ทรัพยากรห้องแลปให้เกิดประโยชน์สูงสุด",
            [
                ("การตรวจสอบโดยผู้มีอำนาจ: ", "อาจารย์ผู้ประสานงานรายวิชาหรือเจ้าหน้าที่ห้องปฏิบัติการตรวจสอบความถูกต้อง"),
                ("จุดตัดสินใจที่ 2 (ผลการพิจารณาอนุมัติหรือไม่?): ", "หากไม่อนุมัติ ระบบจะส่งแจ้งเตือนพร้อมระบุเหตุผลเพื่อให้ผู้ขอนำไปแก้ไขหรือยกเลิก หากอนุมัติ ระบบจะเปลี่ยนสถานะเป็น 'อนุมัติแล้ว' (APPROVED) พร้อมทั้งจองคิวห้องและล็อกยอดพัสดุในระบบทันที")
            ]
        ),
        (
            "ขั้นตอนที่ 4: การจัดเตรียมและส่งมอบอุปกรณ์ (Fulfillment & Check-out)",
            "ขั้นตอนการปฏิบัติการหน้าคลังของเจ้าหน้าที่เพื่อส่งมอบพัสดุอุปกรณ์อย่างถูกต้องและตรวจสอบได้",
            [
                ("การจัดของตามใบงาน: ", "เจ้าหน้าที่พิมพ์ใบจัดพัสดุ และหยิบครุภัณฑ์รายชิ้นพร้อมเวชภัณฑ์ตามจำนวนที่ขอ"),
                ("การสแกน QR Code ตรวจสอบ: ", "เจ้าหน้าที่ใช้เครื่องอ่านหรือโทรศัพท์มือถือสแกน QR Code ประจำอุปกรณ์เพื่อยืนยันหมายเลขเครื่อง"),
                ("การส่งมอบและลงนามรับ: ", "ผู้ขอตรวจรับอุปกรณ์ร่วมกับเจ้าหน้าที่ และกดยืนยันการรับมอบในระบบ สถานะจะปรับเป็น 'อยู่ระหว่างใช้งาน' (IN_USE)")
            ]
        ),
        (
            "ขั้นตอนที่ 5: การฝึกปฏิบัติและการตรวจรับคืน (Practice & Return Inspection)",
            "ขั้นตอนการจัดการหลังเสร็จสิ้นการฝึกปฏิบัติการพยาบาล ซึ่งเป็นจุดสำคัญของการควบคุมทรัพย์สิน",
            [
                ("การฝึกปฏิบัติการ: ", "นักศึกษาใช้ห้องและอุปกรณ์ฝึกทักษะทางการพยาบาลตามระยะเวลาที่กำหนด"),
                ("การนำส่งคืนพัสดุ: ", "เมื่อเสร็จสิ้นการฝึก ผู้ขอนำอุปกรณ์และเวชภัณฑ์คงเหลือมาส่งคืนที่เคาน์เตอร์ห้องแลป"),
                ("การตรวจนับและประเมินสภาพ: ", "เจ้าหน้าที่ตรวจสอบจำนวน ความสะอาด และการทำงานของอุปกรณ์รายชิ้น"),
                ("จุดตัดสินใจที่ 3 (พบการชำรุด เสียหาย หรือสูญหายหรือไม่?): ", "แยกการดำเนินการเป็น 2 แขนงสำคัญ:\n"
                 "  • กรณีพบความเสียหาย/สูญหาย: เจ้าหน้าที่บันทึกประเมินความเสียหาย (Damage Assessment) พร้อมภาพถ่าย -> คำนวณค่าปรับ/ค่าชดใช้ตามระเบียบคณะ -> ปรับสถานะเครื่องเป็นแจ้งซ่อม (MAINTENANCE)\n"
                 "  • กรณีสภาพปกติสมบูรณ์: ยืนยันการรับคืนสำเร็จ -> ปรับสถานะเครื่องกลับเป็นพร้อมใช้งาน (AVAILABLE) -> บันทึกตัดลดยอดเวชภัณฑ์ที่ใช้หมดไปตามจริง (Consumed Stock)")
            ]
        ),
        (
            "ขั้นตอนที่ 6: การบันทึกประวัติ สรุปผล และรายงานผู้บริหาร (Audit & Analytics)",
            "ขั้นตอนสุดท้ายของการปิดวงจรข้อมูลและการวิเคราะห์ผลเชิงยุทธศาสตร์",
            [
                ("การปิดคำร้องสมบูรณ์: ", "ระบบปรับสถานะคำร้องเป็น 'เสร็จสิ้น' (COMPLETED)"),
                ("การบันทึกประวัติโปร่งใส (Audit Trail): ", "บันทึกข้อมูลธุรกรรมทั้งหมด วันที่ เวลา ผู้ดำเนินการ และการเปลี่ยนแปลงสถานะอย่างละเอียด"),
                ("การประมวลผลดัชนีชี้วัด (KPIs): ", "ระบบนำข้อมูลไปประมวลผลเป็นสถิติ เช่น อัตราการใช้ห้องแลป, อุปกรณ์ที่ถูกยืมบ่อย, สถิติการชำรุด และต้นทุนเวชภัณฑ์"),
                ("การแสดงผลบน Executive Dashboard: ", "นำเสนอผลลัพธ์ผ่านแผนภูมิและตัวเลขสรุปเพื่อให้หัวหน้าสาขา คณบดี และผู้บริหารใช้ในการวางแผนจัดซื้อและบริหารงบประมาณ")
            ]
        )
    ]

    for title, desc, items in stages:
        add_heading_styled(title, level=2)
        add_body_p(desc)
        for prefix, txt in items:
            add_body_p(txt, bold_prefix=prefix, bullet=True)

    # Save documents
    doc.save(target_docx)
    doc.save(alias_docx)
    print("Successfully generated Word documents:")
    print(" -", target_docx)
    print(" -", alias_docx)

if __name__ == "__main__":
    create_vertical_flowchart_docx()
